"""Export produces real, openable files from stored Markdown."""
import io, zipfile, pytest
from app import export

SAMPLE = """# Growth Plan

Lagos-based agritech serving **1,200 smallholder farms**.

## Market

Bottom-up: 1,200 farms x N85,000 = N102,000,000 obtainable.

| Segment | Farms | Contract |
|---|---|---|
| Smallholder | 1,200 | N85,000 |
| Cooperative | [NEEDS INPUT: count] | N400,000 |

### Use of funds

- Working capital, 40%
- Hiring, 35%
- Reserve, 25%

1. Close the round
2. Reach 3,000 farms
3. Series A readiness

---

Revenue for FY2025 is [NEEDS INPUT: audited figure].
"""


def test_parser_identifies_blocks():
    kinds = [b["type"] for b in export.parse_blocks(SAMPLE)]
    assert "heading" in kinds and "table" in kinds
    assert "bullet" in kinds and "number" in kinds and "rule" in kinds


def test_table_rows_are_square():
    table = next(b for b in export.parse_blocks(SAMPLE) if b["type"] == "table")
    assert table["head"] == ["Segment", "Farms", "Contract"]
    assert all(len(r) == len(table["head"]) for r in table["rows"])
    assert len(table["rows"]) == 2          # separator row dropped


def test_list_items_parsed():
    blocks = export.parse_blocks(SAMPLE)
    bullets = next(b for b in blocks if b["type"] == "bullet")
    numbers = next(b for b in blocks if b["type"] == "number")
    assert len(bullets["items"]) == 3
    assert numbers["items"][0] == "Close the round"


def test_pdf_is_a_real_pdf():
    data = export.to_pdf(SAMPLE, "Growth Plan", "Prepared for investors")
    assert data[:5] == b"%PDF-"
    assert data.rstrip().endswith(b"%%EOF")
    assert len(data) > 2000


def test_docx_is_a_real_docx():
    data = export.to_docx(SAMPLE, "Growth Plan", "Prepared for investors")
    assert data[:2] == b"PK"
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        assert "word/document.xml" in z.namelist()
        assert z.testzip() is None
        xml = z.read("word/document.xml").decode("utf-8")
    assert "Growth Plan" in xml
    assert "NEEDS INPUT" in xml            # placeholders survive into the file
    assert "**" not in xml                 # markers converted, not literal


def test_docx_roundtrips_through_python_docx():
    from docx import Document
    doc = Document(io.BytesIO(export.to_docx(SAMPLE, "Growth Plan")))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Lagos-based agritech" in text
    assert len(doc.tables) == 1
    assert doc.tables[0].rows[0].cells[0].text == "Segment"


def test_build_refuses_non_document_output():
    with pytest.raises(export.ExportError):
        export.build({"kind": "scorecard", "data": {}}, "pdf", "CV analysis")


def test_build_refuses_unknown_format():
    with pytest.raises(export.ExportError):
        export.build({"kind": "document", "markdown": "# x"}, "rtf", "Plan")


def test_filenames_are_safe_and_meaningful():
    _, _, name = export.build(
        {"kind": "document", "markdown": SAMPLE, "title": "Adaeze Okoro / CV: Product Manager"},
        "pdf", "CV rewrite")
    assert name.endswith(".pdf")
    assert "/" not in name and ":" not in name
    assert "Adaeze" in name
