"""
Sofia — document export (shared service S4).

Turns a completed job's stored Markdown into a PDF or a Word file.

The rule that matters
---------------------
Export formats what is already in `jobs.output_json`. It never calls the
model. A document the user downloads must be byte-for-byte the same content
they reviewed on screen — if the download could differ from the preview, the
preview is worthless and so is the review.

Both formats are built from one block parse, so a heading is a heading in
both and the two never drift apart.

[NEEDS INPUT: ...] placeholders survive into the exported file, highlighted.
That is the whole point of the Evidence Tiering Protocol: the gap has to be
visible in the artefact the user actually sends to someone.
"""

from __future__ import annotations

import io
import re
from datetime import datetime, timezone

# --------------------------------------------------------------------------- #
#  Palette — matches static/css/sofia.css
# --------------------------------------------------------------------------- #
INDIGO = (0x43, 0x38, 0xCA)
INK = (0x14, 0x15, 0x1A)
INK_3 = (0x6B, 0x6F, 0x7B)
BORDER = (0xE4, 0xE6, 0xEC)
AMBER_BG = (0xFF, 0xF8, 0xEB)
AMBER_INK = (0xB5, 0x47, 0x08)

_NEEDS_INPUT = re.compile(r"\[NEEDS INPUT:([^\]]*)\]")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_ITALIC = re.compile(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)")
_CODE = re.compile(r"`([^`]+)`")
_SEPARATOR = re.compile(r"^:?-{2,}:?$")


# --------------------------------------------------------------------------- #
#  Block parser — one pass, both formats consume the result
# --------------------------------------------------------------------------- #
def parse_blocks(markdown: str) -> list[dict]:
    """
    Parse the same Markdown subset that engines.render_markdown accepts.

    Returns a flat list of blocks:
        {"type": "heading", "level": 1-4, "text": str}
        {"type": "para",    "text": str}
        {"type": "bullet",  "items": [str]}
        {"type": "number",  "items": [str]}
        {"type": "table",   "head": [str], "rows": [[str]]}
        {"type": "rule"}
    """
    blocks: list[dict] = []
    lines = (markdown or "").replace("\r\n", "\n").split("\n")
    i = 0

    while i < len(lines):
        raw = lines[i].rstrip()
        line = raw.strip()

        if not line:
            i += 1
            continue

        # --- table -------------------------------------------------------
        if line.startswith("|") and line.endswith("|"):
            rows: list[list[str]] = []
            while i < len(lines):
                candidate = lines[i].strip()
                if not (candidate.startswith("|") and candidate.endswith("|")):
                    break
                cells = [c.strip() for c in candidate.strip("|").split("|")]
                if not all(_SEPARATOR.fullmatch(c) for c in cells if c):
                    rows.append(cells)
                i += 1
            if rows:
                width = max(len(r) for r in rows)
                rows = [r + [""] * (width - len(r)) for r in rows]
                blocks.append({"type": "table", "head": rows[0], "rows": rows[1:]})
            continue

        # --- heading -----------------------------------------------------
        heading = re.match(r"^(#{1,4})\s+(.*)$", line)
        if heading:
            blocks.append({
                "type": "heading",
                "level": len(heading.group(1)),
                "text": heading.group(2).strip(),
            })
            i += 1
            continue

        # --- horizontal rule ---------------------------------------------
        if re.fullmatch(r"(-{3,}|_{3,}|\*{3,})", line):
            blocks.append({"type": "rule"})
            i += 1
            continue

        # --- lists -------------------------------------------------------
        bullet = re.match(r"^[-*+]\s+(.*)$", line)
        number = re.match(r"^\d+[.)]\s+(.*)$", line)
        if bullet or number:
            kind = "bullet" if bullet else "number"
            pattern = r"^[-*+]\s+(.*)$" if bullet else r"^\d+[.)]\s+(.*)$"
            items: list[str] = []
            while i < len(lines):
                match = re.match(pattern, lines[i].strip())
                if not match:
                    # A wrapped continuation line belongs to the item above.
                    following = lines[i].strip()
                    if following and items and not re.match(r"^(#|\||[-*+]\s|\d+[.)]\s)", following):
                        items[-1] += " " + following
                        i += 1
                        continue
                    break
                items.append(match.group(1).strip())
                i += 1
            blocks.append({"type": kind, "items": items})
            continue

        # --- paragraph ---------------------------------------------------
        parts = [line]
        i += 1
        while i < len(lines):
            following = lines[i].strip()
            if not following or re.match(r"^(#{1,4}\s|\||[-*+]\s|\d+[.)]\s|-{3,}|_{3,})", following):
                break
            parts.append(following)
            i += 1
        blocks.append({"type": "para", "text": " ".join(parts)})

    return blocks


def strip_inline(text: str) -> str:
    """Plain text with the Markdown emphasis markers removed."""
    text = _BOLD.sub(r"\1", text)
    text = _ITALIC.sub(r"\1", text)
    text = _CODE.sub(r"\1", text)
    return text


def _filename(base: str, ext: str) -> str:
    """A filename someone can find again in their downloads folder."""
    clean = re.sub(r"[^A-Za-z0-9]+", "_", base or "Sofia_document").strip("_")
    return f"{clean[:70] or 'Sofia_document'}.{ext}"


# --------------------------------------------------------------------------- #
#  PDF
# --------------------------------------------------------------------------- #
def to_pdf(markdown: str, title: str, subtitle: str = "") -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (HRFlowable, ListFlowable, ListItem,
                                    PageBreak, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    def rgb(triple):
        return colors.Color(triple[0] / 255, triple[1] / 255, triple[2] / 255)

    base = getSampleStyleSheet()
    body = ParagraphStyle(
        "SofiaBody", parent=base["BodyText"], fontName="Helvetica",
        fontSize=10, leading=15.5, spaceAfter=7, textColor=rgb(INK),
        alignment=TA_LEFT,
    )
    h_styles = {
        1: ParagraphStyle("SofiaH1", parent=body, fontName="Helvetica-Bold",
                          fontSize=17, leading=22, spaceBefore=16, spaceAfter=9,
                          textColor=rgb(INK)),
        2: ParagraphStyle("SofiaH2", parent=body, fontName="Helvetica-Bold",
                          fontSize=13.5, leading=18, spaceBefore=14, spaceAfter=7,
                          textColor=rgb(INK)),
        3: ParagraphStyle("SofiaH3", parent=body, fontName="Helvetica-Bold",
                          fontSize=11.5, leading=16, spaceBefore=11, spaceAfter=5,
                          textColor=rgb(INDIGO)),
        4: ParagraphStyle("SofiaH4", parent=body, fontName="Helvetica-Oblique",
                          fontSize=10.5, leading=15, spaceBefore=9, spaceAfter=4,
                          textColor=rgb(INK_3)),
    }
    cell = ParagraphStyle("SofiaCell", parent=body, fontSize=8.6, leading=12,
                          spaceAfter=0)
    cell_head = ParagraphStyle("SofiaCellHead", parent=cell,
                               fontName="Helvetica-Bold")

    def inline(text: str) -> str:
        """Markdown emphasis to reportlab's mini-HTML, escaped first."""
        out = (text or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        out = _BOLD.sub(r"<b>\1</b>", out)
        out = _ITALIC.sub(r"<i>\1</i>", out)
        out = _CODE.sub(r'<font face="Courier">\1</font>', out)
        out = _NEEDS_INPUT.sub(
            lambda m: f'<font backColor="#FFF8EB" color="#B54708">'
                      f'<b>NEEDS INPUT:{m.group(1)}</b></font>',
            out,
        )
        return out

    story: list = []
    story.append(Paragraph(inline(title), h_styles[1]))
    if subtitle:
        story.append(Paragraph(
            inline(subtitle),
            ParagraphStyle("SofiaSub", parent=body, fontSize=9.5,
                           textColor=rgb(INK_3), spaceAfter=4),
        ))
    story.append(HRFlowable(width="100%", thickness=0.7, color=rgb(BORDER),
                            spaceBefore=6, spaceAfter=12))

    for block in parse_blocks(markdown):
        kind = block["type"]

        if kind == "heading":
            story.append(Paragraph(inline(block["text"]),
                                   h_styles[min(block["level"], 4)]))

        elif kind == "para":
            story.append(Paragraph(inline(block["text"]), body))

        elif kind in ("bullet", "number"):
            items = [ListItem(Paragraph(inline(t), body), leftIndent=12)
                     for t in block["items"]]
            story.append(ListFlowable(
                items,
                bulletType="bullet" if kind == "bullet" else "1",
                start="•" if kind == "bullet" else None,
                leftIndent=14, bulletFontSize=8,
                bulletColor=rgb(INDIGO), spaceAfter=8,
            ))

        elif kind == "table":
            head = [Paragraph(inline(c), cell_head) for c in block["head"]]
            rows = [[Paragraph(inline(c), cell) for c in r] for r in block["rows"]]
            available = A4[0] - 40 * mm
            columns = max(1, len(block["head"]))
            table = Table([head] + rows, repeatRows=1,
                          colWidths=[available / columns] * columns)
            table.setStyle(TableStyle([
                ("GRID", (0, 0), (-1, -1), 0.5, rgb(BORDER)),
                ("BACKGROUND", (0, 0), (-1, 0), colors.Color(0.97, 0.97, 0.98)),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
                ("RIGHTPADDING", (0, 0), (-1, -1), 6),
                ("TOPPADDING", (0, 0), (-1, -1), 5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
            ]))
            story.append(table)
            story.append(Spacer(1, 9))

        elif kind == "rule":
            story.append(HRFlowable(width="100%", thickness=0.6,
                                    color=rgb(BORDER), spaceBefore=8, spaceAfter=8))

    generated = datetime.now(timezone.utc).strftime("%d %B %Y")

    def furniture(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(rgb(INK_3))
        canvas.drawString(20 * mm, 12 * mm, f"Prepared with Sofia · {generated}")
        canvas.drawRightString(A4[0] - 20 * mm, 12 * mm, str(canvas.getPageNumber()))
        canvas.restoreState()

    buffer = io.BytesIO()
    SimpleDocTemplate(
        buffer, pagesize=A4,
        leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=18 * mm, bottomMargin=20 * mm,
        title=title, author="Sofia",
    ).build(story, onFirstPage=furniture, onLaterPages=furniture)

    return buffer.getvalue()


# --------------------------------------------------------------------------- #
#  Word
# --------------------------------------------------------------------------- #
def to_docx(markdown: str, title: str, subtitle: str = "") -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = Document()

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(7)

    def add_inline(paragraph, text: str) -> None:
        """
        Emit runs so bold, italic and placeholders survive as real Word
        formatting rather than literal asterisks in the document.
        """
        pattern = re.compile(
            r"(\*\*.+?\*\*|(?<!\*)\*(?!\*).+?(?<!\*)\*(?!\*)|`[^`]+`|\[NEEDS INPUT:[^\]]*\])"
        )
        position = 0
        for match in pattern.finditer(text or ""):
            if match.start() > position:
                paragraph.add_run(text[position:match.start()])
            token = match.group(0)
            if token.startswith("**"):
                paragraph.add_run(token[2:-2]).bold = True
            elif token.startswith("`"):
                run = paragraph.add_run(token[1:-1])
                run.font.name = "Consolas"
            elif token.startswith("[NEEDS INPUT:"):
                run = paragraph.add_run(token[1:-1])
                run.bold = True
                run.font.color.rgb = RGBColor(*AMBER_INK)
            else:
                paragraph.add_run(token[1:-1]).italic = True
            position = match.end()
        if position < len(text or ""):
            paragraph.add_run(text[position:])

    heading = document.add_heading(title, level=0)
    for run in heading.runs:
        run.font.color.rgb = RGBColor(*INK)
    if subtitle:
        sub = document.add_paragraph()
        run = sub.add_run(subtitle)
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(*INK_3)

    for block in parse_blocks(markdown):
        kind = block["type"]

        if kind == "heading":
            para = document.add_heading(level=min(block["level"], 4))
            add_inline(para, block["text"])

        elif kind == "para":
            add_inline(document.add_paragraph(), block["text"])

        elif kind in ("bullet", "number"):
            style = "List Bullet" if kind == "bullet" else "List Number"
            for item in block["items"]:
                add_inline(document.add_paragraph(style=style), item)

        elif kind == "table":
            table = document.add_table(rows=1, cols=max(1, len(block["head"])))
            table.style = "Light Grid Accent 1"
            for index, text in enumerate(block["head"]):
                cell_para = table.rows[0].cells[index].paragraphs[0]
                add_inline(cell_para, text)
                for run in cell_para.runs:
                    run.bold = True
            for row in block["rows"]:
                cells = table.add_row().cells
                for index, text in enumerate(row[:len(cells)]):
                    add_inline(cells[index].paragraphs[0], text)
            document.add_paragraph()

        elif kind == "rule":
            para = document.add_paragraph("─" * 40)
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in para.runs:
                run.font.color.rgb = RGBColor(*BORDER)

    footer = document.sections[0].footer.paragraphs[0]
    footer.text = f"Prepared with Sofia · {datetime.now(timezone.utc).strftime('%d %B %Y')}"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor(*INK_3)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --------------------------------------------------------------------------- #
#  Entry point used by the download route
# --------------------------------------------------------------------------- #
FORMATS = {
    "pdf":  ("application/pdf", to_pdf),
    "docx": ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",
             to_docx),
}


class ExportError(Exception):
    pass


def build(payload: dict, fmt: str, tool_name: str) -> tuple[bytes, str, str]:
    """
    Returns (data, mimetype, filename).

    Raises ExportError when this result has nothing exportable — a scorecard
    is a screen artefact, not a document, and offering a download for one
    would produce an empty file.
    """
    if fmt not in FORMATS:
        raise ExportError("That format is not available.")

    markdown = (payload or {}).get("markdown")
    if not markdown:
        raise ExportError("This result cannot be downloaded as a document.")

    title = (payload.get("title") or tool_name or "Sofia document").strip()
    subtitle = (payload.get("subtitle") or "").strip()

    mimetype, renderer = FORMATS[fmt]
    return renderer(markdown, title, subtitle), mimetype, _filename(title, fmt)
